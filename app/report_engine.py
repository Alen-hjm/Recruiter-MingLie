# -*- coding: utf-8 -*-
"""
推荐报告生成引擎。

设计要点（与项目整体一致）：
  1. 零编造 —— 报告里每一个字都来自「简历原文」或「评分阶段已生成的内容」，
     引擎本身不调用 LLM 重新生成任何叙述性文字。
     理由：评分时 LLM 已经基于简历原文分析过并给出了亮点/风险/总结，
     再生成一遍既多花钱，又引入新的幻觉风险。
  2. 可归因 —— 每个字段都记录来源与置信度，缺的一律留空并标「需人工补充」。
  3. 模板驱动 —— 字段清单在 config/report_templates/*.yaml，改模板不用改代码。
  4. 校验先行 —— 生成后检查必填缺失与占位符残留，有问题明确提示，不静默放行。
"""

import os
import re
import json
from datetime import datetime

try:
    import yaml
except ImportError:
    yaml = None

from .resume_parser import parse_resume

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_DIR = os.path.join(BASE_DIR, "config", "report_templates")

PLACEHOLDER_RE = re.compile(r"\{\{[^}]+\}\}")
MISSING_MARK = "[需人工补充]"


# ------------------------------------------------------------------ 模板加载
def load_template(template_id="headhunter_v1"):
    path = os.path.join(TEMPLATE_DIR, "%s.yaml" % template_id)
    if not os.path.exists(path):
        raise FileNotFoundError("模板文件不存在: %s" % path)
    if yaml is None:
        raise RuntimeError("缺少 pyyaml，请先 pip install pyyaml")
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def list_templates():
    if not os.path.isdir(TEMPLATE_DIR):
        return []
    return [f[:-5] for f in os.listdir(TEMPLATE_DIR) if f.endswith(".yaml")]


# ------------------------------------------------------------------ 取值
def safe_resolve(expr, ctx):
    """
    只支持 "a.b.c" 路径取值，不做任何 eval。
    取不到返回空字符串（由调用方决定如何标注）。
    """
    if not expr:
        return ""
    cur = ctx
    for part in str(expr).split("."):
        if isinstance(cur, dict):
            cur = cur.get(part)
        elif hasattr(cur, part):
            cur = getattr(cur, part)
        else:
            return ""
        if cur is None:
            return ""
    return cur


def _apply_format(value, fmt):
    if not fmt:
        return value
    try:
        return str(fmt).replace("{value}", str(value))
    except Exception:
        return value


def _as_text(value, max_chars=0):
    """把值转成文本；列表转成「- 项」形式"""
    if value is None or value == "":
        return ""
    if isinstance(value, (list, tuple)):
        items = [str(v).strip() for v in value if str(v).strip()]
        text = "\n".join("- " + i for i in items)
    else:
        text = str(value).strip()
    if max_chars and len(text) > max_chars:
        text = text[:max_chars] + "…（已截断，原文 %d 字）" % len(text)
    return text


# ------------------------------------------------------------------ 生成
def build_report(resume, score100=None, job=None, template_id="headhunter_v1"):
    """
    生成一份推荐报告。

    参数:
      resume     —— 简历 dict（至少要有 full_text；有 detail_url/name 更好）
      score100   —— 评分结果 dict（score_tasks 里的 score100 字段），可为 None
      job        —— 岗位 dict（可选，用于报告头部）
    返回:
      {
        "template": {...},
        "sections": [ {"key","title","type","content"/"rows","source","confidence"} ],
        "validation": {"ok","missing_required","need_human","placeholder_left"},
        "meta": {...}
      }
    """
    tpl = load_template(template_id)
    score100 = score100 or {}

    parsed = parse_resume(resume or {})
    ctx = {
        "parsed": parsed,
        "score": score100,
        "job": job or {},
        "resume": resume or {},
    }

    sections = []
    missing_required = []
    need_human = []

    for sec in tpl.get("sections", []):
        skey = sec.get("key")
        stype = sec.get("type", "text")
        source = sec.get("source", "code")
        required = bool(sec.get("required"))
        max_chars = int(sec.get("max_chars", 0) or 0)

        out = {
            "key": skey,
            "title": sec.get("title", skey),
            "type": stype,
            "source": source,
            "confidence": "high",
            "note": sec.get("note", ""),
        }

        # ---- 表格类型（基本信息）
        if stype == "table":
            rows = []
            low_conf = False
            for f in sec.get("fields", []):
                raw = safe_resolve(f.get("expr", ""), ctx)
                val = _apply_format(raw, f.get("format"))
                if val in ("", None):
                    val = MISSING_MARK
                    if f.get("required"):
                        missing_required.append({
                            "section": skey, "field": f.get("key"),
                            "label": f.get("label", f.get("key")),
                        })
                    need_human.append(f.get("label", f.get("key")))
                rows.append({
                    "label": f.get("label", f.get("key")),
                    "value": str(val),
                })
                if raw in ("", None):
                    low_conf = True
            out["rows"] = rows
            out["confidence"] = "medium" if low_conf else "high"

        # ---- 文本 / 列表
        else:
            raw = safe_resolve(sec.get("expr", ""), ctx)
            text = _as_text(raw, max_chars)
            if not text:
                text = MISSING_MARK
                if required:
                    missing_required.append({
                        "section": skey, "field": skey, "label": sec.get("title", skey),
                    })
                need_human.append(sec.get("title", skey))
                out["confidence"] = "low"
            out["content"] = text
            # 原文类型顺带记一下切分置信度
            if source == "resume" and skey != "basic":
                conf = (parsed.get("sections", {}).get(skey) or {}).get("confidence")
                if conf:
                    out["confidence"] = conf

        sections.append(out)

    # ---- 占位符残留检查（渲染后统一扫）
    placeholder_left = []
    for s in sections:
        blob = json.dumps(s.get("content") or s.get("rows") or "", ensure_ascii=False)
        placeholder_left.extend(PLACEHOLDER_RE.findall(blob))

    validation = {
        "ok": not missing_required and not placeholder_left,
        "missing_required": missing_required,
        "placeholder_left": sorted(set(placeholder_left)),
        "need_human": sorted(set(need_human)),
        "missing_sections": parsed.get("missing", []),
    }

    return {
        "template": {"id": tpl.get("id"), "name": tpl.get("name"), "version": tpl.get("version")},
        "sections": sections,
        "validation": validation,
        "meta": {
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "parser": parsed.get("meta", {}),
            "job_title": (job or {}).get("title", ""),
            "has_score": bool(score100),
            "score_total": score100.get("total_score", ""),
            "score_pass": score100.get("pass", ""),
        },
    }


# ------------------------------------------------------------------ 渲染 Markdown
def render_markdown(report):
    lines = []
    meta = report.get("meta", {})
    title = report.get("template", {}).get("name", "推荐报告")
    lines.append("# %s" % title)
    head = []
    if meta.get("job_title"):
        head.append("应聘岗位：%s" % meta["job_title"])
    if meta.get("score_total") != "":
        head.append("AI 评分：%s 分（%s）" % (meta["score_total"], "通过" if meta.get("score_pass") else "未通过"))
    head.append("生成时间：%s" % meta.get("generated_at", ""))
    lines.append("> " + " ｜ ".join(head))
    lines.append("")

    for s in report.get("sections", []):
        lines.append("## %s" % s["title"])
        lines.append("")
        if s["type"] == "table":
            lines.append("| 项目 | 内容 |")
            lines.append("| --- | --- |")
            for r in s.get("rows", []):
                lines.append("| %s | %s |" % (r["label"], str(r["value"]).replace("\n", " ")))
        else:
            lines.append(s.get("content", ""))
        lines.append("")

    v = report.get("validation", {})
    if not v.get("ok"):
        lines.append("---")
        lines.append("")
        lines.append("**⚠ 需人工确认**")
        if v.get("missing_required"):
            names = "、".join(sorted(set(m["label"] for m in v["missing_required"])))
            lines.append("- 必填项缺失：%s" % names)
        if v.get("placeholder_left"):
            lines.append("- 存在未替换占位符：%s" % "、".join(v["placeholder_left"]))
        lines.append("- 标注为「%s」的项请人工补填后再发给客户" % MISSING_MARK)
        lines.append("")
    return "\n".join(lines)


# ------------------------------------------------------------------ 渲染 docx
def render_docx(report, out_path):
    """
    用 python-docx 新建文档（不套用用户模板）。
    套用用户 docx 模板留到二期：Word 会把 {{name}} 拆进多个 run，替换容易失败。
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    meta = report.get("meta", {})

    title = doc.add_heading(report.get("template", {}).get("name", "推荐报告"), level=0)
    sub = []
    if meta.get("job_title"):
        sub.append("应聘岗位：%s" % meta["job_title"])
    if meta.get("score_total") != "":
        sub.append("AI 评分：%s 分（%s）" % (meta["score_total"], "通过" if meta.get("score_pass") else "未通过"))
    sub.append("生成时间：%s" % meta.get("generated_at", ""))
    p = doc.add_paragraph(" ｜ ".join(sub))

    for s in report.get("sections", []):
        doc.add_heading(s["title"], level=1)
        if s["type"] == "table":
            rows = s.get("rows", [])
            table = doc.add_table(rows=0, cols=2)
            table.style = "Light Grid Accent 1"
            for r in rows:
                cells = table.add_row().cells
                cells[0].text = r["label"]
                cells[1].text = str(r["value"])
        else:
            content = s.get("content", "")
            if "\n- " in content or content.startswith("- "):
                for line in content.split("\n"):
                    line = line.strip()
                    if line.startswith("- "):
                        doc.add_paragraph(line[2:], style="List Bullet")
                    elif line:
                        doc.add_paragraph(line)
            else:
                doc.add_paragraph(content)

    v = report.get("validation", {})
    if not v.get("ok"):
        doc.add_heading("需人工确认", level=1)
        if v.get("missing_required"):
            names = "、".join(sorted(set(m["label"] for m in v["missing_required"])))
            doc.add_paragraph("必填项缺失：%s" % names)
        if v.get("placeholder_left"):
            doc.add_paragraph("未替换占位符：%s" % "、".join(v["placeholder_left"]))
        doc.add_paragraph("标注为「%s」的项请人工补填后再发给客户。" % MISSING_MARK)

    doc.save(out_path)
    return out_path
